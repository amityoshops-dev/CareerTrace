"""CareerTrace Pro: multi-portal search, resume builder, Excel tracker, signed emails, LinkedIn drafts.
Loaded by careertrace_ui.py via careertrace_pro.install(app, globals()). A failure here never stops the app."""
import os, re, io, csv, json, base64, uuid, ssl, smtplib, threading, time, urllib.request, urllib.parse, urllib.error, datetime as dt
from concurrent.futures import ThreadPoolExecutor
from email.message import EmailMessage
from pathlib import Path
from typing import List
from urllib.parse import urlparse
from dotenv import load_dotenv
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from pydantic import BaseModel

load_dotenv()
load_dotenv(".env.pro")
HERE = Path(__file__).parent
STATIC = HERE / "pro_static"
G = {}
STATE = {"conns": [], "msgs": {}}
AG = None
RESUME_OUT = "output/Amit_Mankar_Resume_Updated.docx"
TRACKER = "exports/CareerTrace_Tracker.xlsx"

DEFAULT_PROFILE = {
    "name": "Amit Mankar", "email": "avmankar001@gmail.com", "phone": "9922080307",
    "linkedin": "https://www.linkedin.com/in/amit-mankar5",
    "blog": "https://substack.com/@intentcuriositysphere",
    "alumni_terms": "Kotak, Aditya Birla, University of Mumbai, Amravati",
    "projects": [
        {"name": "TBG CORE // Institutional Transaction Banking Platform", "url": "https://tbg-engine.onrender.com/", "summary": ""},
        {"name": "NPCI Multi-Rail Gateway Cockpit & Architecture Blueprint", "url": "https://npci-b2b-engine.onrender.com/", "summary": ""},
        {"name": "Institutional Fintech Suite - Business Banking Suite", "url": "https://institutional-fintech-suite.onrender.com/", "summary": ""},
    ],
}


def env(k):
    return os.environ.get(k, "")


def profile():
    os.makedirs("data", exist_ok=True)
    if not os.path.exists("data/profile.json"):
        json.dump(DEFAULT_PROFILE, open("data/profile.json", "w"), indent=1)
    P = json.load(open("data/profile.json"))
    P.setdefault("alumni_terms", DEFAULT_PROFILE["alumni_terms"])
    return P


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9+#]+", " ", (s or "").lower())).strip()


def has(text, kw):
    k = norm(kw)
    return bool(k) and re.search(r"(?<![a-z0-9])" + re.escape(k) + r"(?![a-z0-9])", text) is not None


def split(s):
    return [x.strip() for x in re.split(r"[,\n;]", s or "") if x.strip()]


# ---------------- API key pool (auto-failover), settings, AI, exports ----------------
KEYS_FILE = "data/keys.json"
SETTINGS_FILE = "data/settings.json"
_lock = threading.RLock()
PROVIDERS = {
    "crustdata": {"label": "Crustdata", "kind": "Job listings (primary source)", "env": ["CRUSTDATA_API_KEY"], "hint": "Crustdata API key"},
    "jsearch": {"label": "JSearch (RapidAPI)", "kind": "LinkedIn, Naukri, Indeed, Glassdoor, Apna via Google Jobs", "env": ["RAPIDAPI_KEY"], "hint": "RapidAPI key"},
    "adzuna": {"label": "Adzuna", "kind": "Job listings", "env": [], "hint": "Format APP_ID:APP_KEY"},
    "jooble": {"label": "Jooble", "kind": "Job listings", "env": ["JOOBLE_KEY"], "hint": "Jooble API key"},
    "gemini": {"label": "Gemini", "kind": "AI: resume matching, emails, LinkedIn drafts", "env": ["GOOGLE_API_KEY"], "hint": "Starts with AIza"},
    "gmail": {"label": "Gmail app password", "kind": "Send emails from the app", "env": ["GMAIL_APP_PASSWORD"], "hint": "16-character Google app password"},
}
DEFAULT_SEARCH = {
    "titles": "Transaction Banking, Cash Management, Cash Management Services, Corporate Cash Management, CMS Product",
    "location": "Pune, Mumbai", "keywords": "cash management, collections, payments", "must": "", "exclude": "",
    "days": 0, "strict_location": True, "min_score": 40, "max_jobs": 25, "max_queries": 10,
    "dry_run": False, "only_live": True, "exact_title": True, "resume_path": "",
}


class KeyBad(Exception):
    pass


def mask(k):
    return k[:4] + "..." + k[-4:] if len(k) > 12 else "****"


def _ksave(d):
    os.makedirs("data", exist_ok=True)
    with open(KEYS_FILE, "w") as f:
        json.dump(d, f, indent=1)
    try:
        os.chmod(KEYS_FILE, 0o600)
    except Exception:
        pass


def _kload():
    d = _load(KEYS_FILE, {})
    changed = False
    for name, meta in PROVIDERS.items():
        d.setdefault(name, [])
        if not d[name]:
            vals = [env(e) for e in meta["env"] if env(e)]
            if name == "adzuna" and env("ADZUNA_APP_ID") and env("ADZUNA_APP_KEY"):
                vals = [env("ADZUNA_APP_ID") + ":" + env("ADZUNA_APP_KEY")]
            for v in vals:
                d[name].append({"id": uuid.uuid4().hex[:8], "key": v.strip(), "label": "from .env", "off": False, "status": "", "used": ""})
                changed = True
    if changed:
        _ksave(d)
    return d


def pool(name):
    """Usable keys for a provider, healthy ones first (failed keys are retried last, never stuck)."""
    with _lock:
        ks = [e for e in _kload().get(name, []) if not e.get("off") and e.get("key")]
    return sorted(ks, key=lambda e: bool(e.get("status")) and e["status"] != "ok")


def mark(name, key, status):
    with _lock:
        d = _kload()
        for e in d.get(name, []):
            if e["key"] == key:
                e["status"] = status or "ok"
                e["used"] = _now()
        _ksave(d)


def with_keys(name, fn):
    """Run fn(key) with each key in turn; a rejected or exhausted key is flagged and the next one is used."""
    ks = pool(name)
    if not ks:
        raise NoKey("no key - add one in Settings > API keys")
    last = None
    for e in ks:
        try:
            r = fn(e["key"])
            if e.get("status") != "ok":
                mark(name, e["key"], "")
            return r
        except KeyBad as ex:
            mark(name, e["key"], str(ex)[:80])
            last = ex
    raise RuntimeError(f"all {len(ks)} key(s) failed - last: {last}. Add a new key in Settings > API keys")


def settings():
    s = _load(SETTINGS_FILE, {})
    s.setdefault("providers_enabled", {})
    s.setdefault("autorun", True)
    s.setdefault("search", dict(DEFAULT_SEARCH))
    return s


def save_settings(s):
    os.makedirs("data", exist_ok=True)
    json.dump(s, open(SETTINGS_FILE, "w"), indent=1)


def enabled(name):
    return settings()["providers_enabled"].get(name, True)


def gemini_call(prompt):
    model = env("GEMINI_MODEL") or "gemini-2.5-flash"

    def go(k):
        body = json.dumps({"contents": [{"parts": [{"text": prompt}]}]}).encode()
        data = _http(f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={k}",
                     {"Content-Type": "application/json"}, body)
        try:
            return data["candidates"][0]["content"]["parts"][0]["text"]
        except Exception:
            raise RuntimeError("Gemini returned no text")
    return with_keys("gemini", go)


def llm_text(prompt):
    return gemini_call(prompt).strip()


def llm_json(prompt):
    t = gemini_call(prompt).strip()
    t = re.sub(r"^```(?:json)?|```$", "", t, flags=re.M).strip()
    return json.loads(t)


def test_key(name, key):
    try:
        if name == "gemini":
            _http("https://generativelanguage.googleapis.com/v1beta/models?key=" + key)
        elif name == "crustdata":
            import requests
            rr = requests.post(AG.CRUSTDATA_URL, timeout=30, headers={"Authorization": f"Bearer {key}", "x-api-version": AG.CRUSTDATA_VERSION},
                               json={"filters": {"field": "job_details.title", "type": "(.)", "value": "Cash Management"}, "limit": 1})
            if rr.status_code != 200:
                return False, f"HTTP {rr.status_code} {rr.text[:100]}"
        elif name == "jsearch":
            p_jsearch("cash management", "Pune", {"days": 0}, key)
        elif name == "adzuna":
            p_adzuna("cash management", "India", {}, key)
        elif name == "jooble":
            p_jooble("cash management", "India", {}, key)
        elif name == "gmail":
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as sm:
                sm.login(profile()["email"], key)
        return True, "Key works"
    except KeyBad as e:
        return False, str(e)
    except Exception as e:
        return False, f"{type(e).__name__}: {str(e)[:140]}"


def export_jobs(res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    os.makedirs("exports", exist_ok=True)
    cols = ["Score", "Title", "Company", "Location", "Portal", "Posted", "Availability", "Resume fit %", "Matched keywords",
            "Skill gaps", "Referral contacts", "Job URL", "Find connections on LinkedIn"]
    rows = []
    for j in res.get("jobs", []):
        li = j.get("li_links") or {}
        rows.append([j.get("score", ""), j.get("title", ""), j.get("company", ""), j.get("location", ""), j.get("portal", ""),
                     j.get("posted", ""), {"live": "Looks open", "unknown": "Unverified"}.get(j.get("live"), j.get("live", "")),
                     "" if j.get("resume_pct") is None else j["resume_pct"], ", ".join(j.get("matched", [])), ", ".join(j.get("gaps", [])),
                     "; ".join(f"{m['person']} ({m['why']})" for m in j.get("referrals", [])), j.get("url", ""),
                     next(iter(li.values()), "")])
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLM", (8, 38, 26, 24, 14, 12, 14, 12, 30, 30, 40, 50, 50)):
        ws.column_dimensions[col].width = w
    names = []
    for nm in (f"jobs_{stamp}", "jobs_latest"):
        wb.save(f"exports/{nm}.xlsx")
        with open(f"exports/{nm}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(cols)
            w.writerows(rows)
        names += [nm + ".xlsx", nm + ".csv"]
    return names


def do_search(cfg):
    full = dict(DEFAULT_SEARCH)
    full.update(cfg)
    res = run_search(full)
    try:
        res["files"] = export_jobs(res)
    except Exception as e:
        res["files"] = []
        res["export_error"] = str(e)
    return res


def autorun():
    time.sleep(5)
    try:
        s = settings()
        if not s.get("autorun", True):
            return
        cfg = dict(DEFAULT_SEARCH)
        cfg.update(s.get("search", {}))
        cfg["dry_run"] = False
        res = do_search(cfg)
        STATE["last"] = {"time": _now(), "count": len(res["jobs"]), "files": res.get("files", []), "note": res.get("note", "")}
        print(f"CareerTrace auto-run done: {len(res['jobs'])} jobs saved to exports/")
    except Exception as e:
        STATE["last"] = {"time": _now(), "count": 0, "files": [], "note": f"Auto-run failed: {e}"}
        print("CareerTrace auto-run failed:", e)


# ---------------- signature on every email ----------------
def footer():
    P = profile()
    pr = "\n".join(f"  - {p['name']}: {p['url']}" for p in P.get("projects", []) if p.get("url"))
    return (f"{P['name']}\n{P['email']} | {P['phone']}\nLinkedIn: {P['linkedin']}\n"
            f"Blog (Intent Curiosity Sphere): {P['blog']}\n\nSelected projects:\n{pr}")


def sign(text):
    P = profile()
    u = next((p["url"] for p in P.get("projects", []) if p.get("url")), "")
    if u and u in text:
        return text.replace("[Your Name]", P["name"])
    t = text.replace("[Your Name]", "").rstrip()
    return t + ("\n" if t.endswith(",") else "\n\n") + footer()


# ---------------- job sources ----------------
PORTALS = {"linkedin.com": "LinkedIn", "naukri.com": "Naukri", "wellfound.com": "Wellfound", "angel.co": "Wellfound",
           "apna.co": "Apna", "indeed.com": "Indeed", "foundit.in": "Foundit", "monsterindia": "Foundit",
           "glassdoor": "Glassdoor", "instahyre.com": "Instahyre", "cutshort.io": "Cutshort", "shine.com": "Shine",
           "timesjobs.com": "TimesJobs", "iimjobs": "iimjobs", "hirist": "Hirist"}
ALIAS = {"mumbai": ["mumbai", "navi mumbai", "thane", "bombay", "mahape", "andheri", "powai", "bkc", "vashi"],
         "pune": ["pune", "pimpri", "chinchwad", "hinjewadi", "kharadi", "baner"],
         "bengaluru": ["bengaluru", "bangalore"], "gurgaon": ["gurgaon", "gurugram"],
         "delhi": ["delhi", "new delhi", "noida", "ncr"], "hyderabad": ["hyderabad", "secunderabad"]}


class NoKey(Exception):
    pass


def portal_of(url, hint=""):
    h = urlparse(url or "").netloc.lower()
    for k, v in PORTALS.items():
        if k in h:
            return v
    hn = norm(hint)
    for k, v in PORTALS.items():
        if hn and norm(k.split(".")[0]) in hn:
            return v
    return (hint or h.replace("www.", "") or "Other")[:24]


def _http(url, headers=None, data=None):
    req = urllib.request.Request(url, data=data, headers=headers or {"User-Agent": "CareerTrace/2"})
    try:
        with urllib.request.urlopen(req, timeout=25) as r:
            return json.loads(r.read().decode("utf-8", "ignore"))
    except urllib.error.HTTPError as e:
        body = ""
        try:
            body = e.read().decode("utf-8", "ignore")[:300].lower()
        except Exception:
            pass
        if e.code in (401, 402, 403, 429) or (e.code == 400 and "api key" in body):
            raise KeyBad(f"HTTP {e.code}: key rejected or quota used up")
        raise


def _job(title, company, location, url, desc, posted, hint, src):
    return {"title": title or "", "company": company or "", "location": location or "", "url": url or "",
            "description": desc or "", "posted": (posted or "")[:10], "portal": portal_of(url, hint),
            "workplace_type": "", "source": src}


def p_jsearch(q, loc, cfg, k):
    d = cfg.get("days") or 0
    dp = "all" if not d else "today" if d <= 1 else "3days" if d <= 3 else "week" if d <= 7 else "month"
    qs = urllib.parse.urlencode({"query": f"{q} in {loc}", "page": 1, "num_pages": 2, "country": "in", "date_posted": dp})
    data = _http("https://jsearch.p.rapidapi.com/search?" + qs,
                 {"X-RapidAPI-Key": k, "X-RapidAPI-Host": "jsearch.p.rapidapi.com"})
    return [_job(x.get("job_title"), x.get("employer_name"),
                 ", ".join(filter(None, [x.get("job_city"), x.get("job_state"), x.get("job_country")])),
                 x.get("job_apply_link"), x.get("job_description"), x.get("job_posted_at_datetime_utc"),
                 x.get("job_publisher"), "JSearch") for x in data.get("data", [])]


def p_adzuna(q, loc, cfg, key):
    i, _, k = key.partition(":")
    p = {"app_id": i, "app_key": k, "results_per_page": 30, "what": q, "content-type": "application/json"}
    if norm(loc) not in ("india", "all india", ""):
        p["where"] = loc
    if cfg.get("days"):
        p["max_days_old"] = cfg["days"]
    data = _http("https://api.adzuna.com/v1/api/jobs/in/search/1?" + urllib.parse.urlencode(p))
    return [_job(x.get("title"), (x.get("company") or {}).get("display_name"),
                 (x.get("location") or {}).get("display_name"), x.get("redirect_url"),
                 x.get("description"), x.get("created"), "", "Adzuna") for x in data.get("results", [])]


def p_jooble(q, loc, cfg, k):
    body = json.dumps({"keywords": q, "location": "" if norm(loc) in ("india", "all india") else loc, "page": 1}).encode()
    data = _http("https://jooble.org/api/" + k, {"Content-Type": "application/json"}, body)
    return [_job(x.get("title"), x.get("company"), x.get("location"), x.get("link"),
                 re.sub("<[^>]+>", " ", x.get("snippet") or ""), x.get("updated"), x.get("source"), "Jooble")
            for x in data.get("jobs", [])]


def crust_all(titles, locs, kws, limit):
    """Your existing Crustdata search, once per location; rotates to the next key if one is exhausted."""
    def go(k):
        os.environ["CRUSTDATA_API_KEY"] = k
        out, errs = [], []
        for l in locs:
            try:
                jobs, _ = G["crust"](titles, l, kws, limit)
                out += jobs
            except HTTPException as e:
                d = str(e.detail)
                if re.search(r"Crustdata (401|402|403|429)", d):
                    raise KeyBad(d[:80])
                errs.append(f"{l}: {d[:100]}")
        if errs and not out:
            raise RuntimeError("; ".join(errs))
        return out
    out = with_keys("crustdata", go)
    for j in out:
        j["portal"] = j.get("portal") or portal_of(j.get("url", ""), "LinkedIn")
    return out


SAMPLES = [
    _job("Cash Management Sales - Transaction Banking", "Sample Bank", "Mumbai, Maharashtra", "https://www.linkedin.com/jobs/view/1",
         "Own corporate cash management and collections portfolio, payments, virtual accounts, H2H. " * 4, dt.date.today().isoformat(), "", "sample"),
    _job("Relationship Manager - Cash Management", "Sample Bank 2", "Bengaluru", "https://wellfound.com/jobs/3",
         "Manage client relationships for cash management services and payments. " * 5, "", "", "sample"),
    _job("Software Engineer - Java", "Sample Tech", "Pune", "https://apna.co/job/4", "Build microservices in Java. " * 15, "", "", "sample"),
    _job("Transaction Banking Analyst", "Sample Bank 3", "Pune", "https://in.indeed.com/viewjob?jk=5",
         "Analyse collections and payments data for corporate clients; cash management reporting. " * 5, "", "", "sample"),
]


DEAD = ("no longer accepting applications", "no longer available", "job has expired", "this job has expired", "job expired",
        "position has been filled", "position is no longer", "job is closed", "job has been closed", "posting has expired",
        "no longer open", "job not found", "page not found", "vacancy has been closed", "this job is no longer",
        "listing has expired", "job is no longer", "vacancy is closed", "expired job")


def check_live(url):
    """live = page loads with no closed notice; closed = 404/410 or 'no longer available' text; unknown = blocked/unreachable."""
    if not url or "example.com" in url:
        return "unknown"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36",
                                                   "Accept-Language": "en"})
        with urllib.request.urlopen(req, timeout=8) as r:
            html = r.read(200000).decode("utf-8", "ignore").lower()
        return "closed" if any(x in html for x in DEAD) else "live"
    except urllib.error.HTTPError as e:
        return "closed" if e.code in (404, 410) else "unknown"
    except Exception:
        return "unknown"


def loc_ok(jloc, locs):
    ln = [norm(l) for l in locs]
    if not ln or any(l in ("india", "all india", "anywhere", "remote") for l in ln):
        return True
    j = norm(jloc)
    if not j or "remote" in j:
        return True
    return any(any(a in j for a in ALIAS.get(l, [l])) for l in ln)


def score_job(j, cfg, dropped):
    titles, anyk = split(cfg["titles"]), split(cfg["keywords"])
    must, excl = split(cfg["must"]), split(cfg["exclude"])
    locs = split(cfg["location"]) or ["India"]
    desc = j.get("description") or ""
    ttl = norm(j.get("title"))
    text = norm((j.get("title") or "") + " " + desc)
    full = len(desc) >= 300
    if any(has(text, e) for e in excl):
        dropped["excluded"] += 1
        return None
    if cfg["strict_location"] and not loc_ok(j.get("location", ""), locs):
        dropped["location"] += 1
        return None
    days = cfg["days"] or (60 if cfg.get("only_live", True) else 0)
    if days and j.get("posted"):
        try:
            if (dt.date.today() - dt.date.fromisoformat(j["posted"])).days > days:
                dropped["too_old"] += 1
                return None
        except ValueError:
            pass
    mh = [k for k in must if has(text, k)]
    ah = [k for k in anyk if has(text, k)]
    if must and len(mh) < len(must) and full:
        dropped["missing_must_have"] += 1
        return None
    if anyk and not ah and full:
        dropped["no_keyword_match"] += 1
        return None
    ts = 1.0 if not titles else 0.0
    for t in titles:
        w = [x for x in norm(t).split() if len(x) > 2] or norm(t).split()
        if w:
            ts = max(ts, 1.0 if has(ttl, t) else 0.8 * sum(has(ttl, x) for x in w) / len(w))
    if cfg.get("exact_title") and titles and ts < 1.0:
        dropped["title_mismatch"] += 1
        return None
    kws = must + anyk
    matched = mh + [k for k in ah if k not in mh]
    cov = len(matched) / len(kws) if kws else 1.0
    sc = round(50 * ts + 40 * cov + (10 if j.get("location") else 5))
    return dict(j, portal=j.get("portal") or portal_of(j.get("url", "")), kw_score=sc, score=sc,
                matched=matched, partial=not full, supported=[], gaps=[], resume_pct=None)


def resume_text(path=""):
    try:
        return G["read_resume"](path or (RESUME_OUT if os.path.exists(RESUME_OUT) else G["find_resume"]()))
    except Exception:
        return ""


def resume_match(j, rtext):
    try:
        req = AG.llm_json("List up to 12 key skills, domain terms, products or qualifications required in this job text. "
                          "Return ONLY a JSON array of short strings.\n\n" + ((j.get("description") or "")[:4000] or j["title"]))
    except Exception:
        req = []
    if req and rtext:
        sup, gaps, pct = G["match"](rtext, req)
        j.update(supported=sup, gaps=gaps, resume_pct=pct, score=round(0.6 * j["kw_score"] + 0.4 * pct))


def portal_links(titles, locs):
    out = []
    for t in titles[:5]:
        for l in locs[:3]:
            q, lq = urllib.parse.quote_plus(t), urllib.parse.quote_plus(l)
            slug = re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")
            ls = re.sub(r"[^a-z0-9]+", "-", l.lower()).strip("-")
            india = norm(l) in ("india", "all india")
            g = lambda d: "https://www.google.com/search?q=" + urllib.parse.quote_plus(f'site:{d} "{t}" {l}')
            out.append({"title": t, "location": l, "links": {
                "LinkedIn": f"https://www.linkedin.com/jobs/search/?keywords={q}&location={lq}",
                "Naukri": f"https://www.naukri.com/{slug}-jobs" + ("" if india else f"-in-{ls}"),
                "Indeed": f"https://in.indeed.com/jobs?q={q}&l={lq}",
                "Foundit": f"https://www.foundit.in/srp/results?query={q}&locations={lq}",
                "Wellfound": g("wellfound.com/jobs"), "Apna": g("apna.co"),
                "Instahyre": g("instahyre.com"), "Cutshort": g("cutshort.io")}})
    return out


def run_search(cfg):
    titles = split(cfg["titles"])
    locs = split(cfg["location"]) or ["India"]
    kws = split(cfg["keywords"]) + split(cfg["must"])
    status, raw = {}, []
    if cfg["dry_run"]:
        raw = list(G.get("DEMO", [])) + SAMPLES
        status["Sample data"] = {"ok": True, "count": len(raw), "error": ""}
    else:
        pairs = [(t, l) for t in titles for l in locs][:cfg["max_queries"]]
        calls = []
        if enabled("crustdata"):
            calls.append(("Crustdata", lambda: crust_all(titles, locs, kws, min(cfg["max_jobs"] * 2, 50))))
        for kn, nm, fn in (("jsearch", "JSearch (LinkedIn/Naukri/Indeed/Glassdoor/Apna)", p_jsearch),
                           ("adzuna", "Adzuna", p_adzuna), ("jooble", "Jooble", p_jooble)):
            if enabled(kn):
                calls.append((nm, (lambda fn=fn, kn=kn: with_keys(kn, lambda k: [x for q, l in pairs for x in fn(q, l, cfg, k)]))))

        def run(c):
            try:
                return c[0], c[1](), ""
            except NoKey as e:
                return c[0], [], str(e)
            except Exception as e:
                return c[0], [], f"{type(e).__name__}: {str(e)[:140]}"

        with ThreadPoolExecutor(4) as ex:
            for nm, r, err in ex.map(run, calls):
                status[nm] = {"ok": not err, "count": len(r), "error": err}
                raw += r
    dropped = {k: 0 for k in ("excluded", "location", "too_old", "closed", "title_mismatch", "missing_must_have", "no_keyword_match", "low_score")}
    best = {}
    for j in raw:
        sj = score_job(j, cfg, dropped)
        if not sj:
            continue
        key = norm(sj["title"]) + "|" + norm(sj["company"])
        if key not in best or len(sj.get("description", "")) > len(best[key].get("description", "")):
            best[key] = sj
    only_live = cfg.get("only_live", True)
    jobs = sorted(best.values(), key=lambda x: -x["score"])[: cfg["max_jobs"] * (2 if only_live else 1)]
    for j in jobs:
        j["live"] = "unknown"
    if only_live and not cfg["dry_run"]:
        with ThreadPoolExecutor(8) as ex:
            for j, st in zip(jobs, ex.map(lambda j: check_live(j.get("url", "")), jobs)):
                j["live"] = st
        n0 = len(jobs)
        jobs = [j for j in jobs if j["live"] != "closed"]
        dropped["closed"] += n0 - len(jobs)
    jobs = jobs[: cfg["max_jobs"]]
    rpath = cfg.get("resume_path") or (RESUME_OUT if os.path.exists(RESUME_OUT) else G["find_resume"]())
    rt = resume_text(rpath)
    with ThreadPoolExecutor(4) as ex:
        list(ex.map(lambda j: resume_match(j, rt), jobs[:12]))
    keep = []
    for j in jobs:
        if j["score"] < cfg["min_score"]:
            dropped["low_score"] += 1
        else:
            keep.append(j)
    keep.sort(key=lambda x: -x["score"])
    terms = split(profile().get("alumni_terms", ""))
    for j in keep:
        j["description"] = (j.get("description") or "")[:500]
        ms = find_matches(STATE["conns"], j.get("company", ""), j.get("title", ""), terms)
        j["referrals"] = [{"person": f"{p['first']} {p['last']}".strip(), "position": p["position"], "company": p["company"],
                           "profile": p["url"], "why": p["why"]} for p in ms]
        cq = urllib.parse.quote_plus(j.get("company", ""))
        rq = urllib.parse.quote_plus(j.get("company", "") + " recruiter")
        net = "&network=%5B%22F%22%2C%22S%22%5D&origin=FACETED_SEARCH"
        j["li_links"] = {"My connections + 2nd-degree at company": f"https://www.linkedin.com/search/results/people/?keywords={cq}{net}",
                         "Recruiters at company": f"https://www.linkedin.com/search/results/people/?keywords={rq}{net}"}
    note = "" if any(s["ok"] for s in status.values()) or cfg["dry_run"] else \
        "No job source responded. Add or replace API keys in Settings, or use the portal links below."
    return {"jobs": keep, "providers": status, "dropped": dropped, "raw_count": len(raw), "note": note,
            "portal_links": portal_links(titles, locs), "resume": rpath, "resume_read": bool(rt)}


# ---------------- Excel tracker ----------------
STATUSES = ["Saved", "Applied", "Referral Requested", "Recruiter Contacted", "Interview Scheduled",
            "Interview Done", "Offer", "Rejected", "On Hold", "Withdrawn"]
APP = ["Application ID", "Date Added", "Job Title", "Company", "Location", "Portal", "Job URL", "Status", "Date Applied",
       "Match Score", "Matched Keywords", "Keyword Gaps (not in resume)", "Contact Name", "Contact Email",
       "Follow-up Date", "Notes", "Last Updated"]
LOG = ["Timestamp", "Application ID", "Company", "Job Title", "Old Status", "New Status", "Note"]
OUT = ["Timestamp", "Application ID", "Company", "Person", "Relationship", "Channel", "Message", "Status"]


def _now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def _wb():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if os.path.exists(TRACKER):
        return load_workbook(TRACKER)
    os.makedirs("exports", exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    def sh(name, cols, w=None):
        ws = wb.create_sheet(name)
        ws.append(cols)
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            ws.column_dimensions[c.column_letter].width = (w or {}).get(c.value, 18)
        return ws

    a = sh("Applications", APP, {"Job Title": 32, "Company": 24, "Job URL": 40, "Notes": 40,
                                 "Matched Keywords": 30, "Keyword Gaps (not in resume)": 30})
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES))
    a.add_data_validation(dv)
    dv.add("H2:H2000")
    sh("Status Log", LOG, {"Job Title": 32, "Company": 24, "Note": 40})
    sh("Outreach", OUT, {"Message": 80, "Person": 24, "Company": 24})
    d = wb.create_sheet("Dashboard", 0)
    d["A1"] = "CareerTrace Dashboard"
    d["A1"].font = Font(bold=True, size=16, color="1F3864")
    for i, s in enumerate(STATUSES, 3):
        d.cell(i, 1, s)
        d.cell(i, 2, f"=COUNTIF(Applications!H:H,A{i})")
    d.cell(len(STATUSES) + 3, 1, "Total")
    d.cell(len(STATUSES) + 3, 2, "=COUNTA(Applications!A:A)-1")
    d.column_dimensions["A"].width = 26
    wb.save(TRACKER)
    return wb


def list_apps():
    ws = _wb()["Applications"]
    return [dict(zip(APP, ["" if v is None else str(v) for v in r])) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def add_app(job, status="Saved"):
    for r in list_apps():
        if (job.get("url") and r["Job URL"] == job["url"]) or \
                norm(r["Job Title"]) + norm(r["Company"]) == norm(job.get("title", "")) + norm(job.get("company", "")):
            return r
    wb = _wb()
    aid = "CT-" + dt.date.today().strftime("%y%m%d") + "-" + uuid.uuid4().hex[:4].upper()
    row = {"Application ID": aid, "Date Added": _now(), "Job Title": job.get("title", ""), "Company": job.get("company", ""),
           "Location": job.get("location", ""), "Portal": job.get("portal", ""), "Job URL": job.get("url", ""),
           "Status": status, "Date Applied": _now() if status == "Applied" else "", "Match Score": job.get("score", ""),
           "Matched Keywords": ", ".join(job.get("matched", [])),
           "Keyword Gaps (not in resume)": ", ".join(job.get("gaps", [])), "Last Updated": _now()}
    wb["Applications"].append([row.get(c, "") for c in APP])
    wb["Status Log"].append([_now(), aid, row["Company"], row["Job Title"], "", status, "Created"])
    wb.save(TRACKER)
    return row


def set_status(aid, status, note="", **fields):
    wb = _wb()
    ws = wb["Applications"]
    h = {c.value: i for i, c in enumerate(ws[1])}
    for r in ws.iter_rows(min_row=2):
        if r[0].value == aid:
            old = r[h["Status"]].value
            r[h["Status"]].value = status
            r[h["Last Updated"]].value = _now()
            if status == "Applied" and not r[h["Date Applied"]].value:
                r[h["Date Applied"]].value = _now()
            if note:
                r[h["Notes"]].value = ((r[h["Notes"]].value or "") + " | " + note).strip(" |")
            for k, v in fields.items():
                if k in h and v:
                    r[h[k]].value = v
            wb["Status Log"].append([_now(), aid, r[h["Company"]].value, r[h["Job Title"]].value, old, status, note])
            wb.save(TRACKER)
            return True
    return False


def log_outreach(aid, company, person, msg):
    wb = _wb()
    wb["Outreach"].append([_now(), aid, company, person, "1st-degree", "LinkedIn", msg, "Drafted"])
    wb.save(TRACKER)


# ---------------- LinkedIn network ----------------
def parse_connections(t):
    lines = t.lstrip("\ufeff").splitlines()
    i = next((k for k, l in enumerate(lines) if l.lower().startswith("first name")), 0)
    return [{"first": r.get("First Name", "").strip(), "last": r.get("Last Name", "").strip(), "url": r.get("URL", "").strip(),
             "company": r.get("Company", "").strip(), "position": r.get("Position", "").strip(),
             "connected_on": r.get("Connected On", "").strip()} for r in csv.DictReader(io.StringIO("\n".join(lines[i:])))]


def parse_messages(t, me="amit-mankar5"):
    out = {}
    for r in csv.DictReader(io.StringIO(t.lstrip("\ufeff"))):
        r = {(k or "").upper(): v for k, v in r.items()}
        d = (r.get("DATE") or "")[:10]
        blob = (r.get("SENDER PROFILE URL") or "") + " " + (r.get("RECIPIENT PROFILE URLS") or "")
        for u in re.findall(r"https?://[^\s,;\"]*linkedin\.com/in/[^\s,;\"/]+", blob):
            u = u.lower().rstrip("/")
            if me not in u and d > out.get(u, ""):
                out[u] = d
    return out


def _fmt(d):
    try:
        return dt.datetime.strptime(d[:10], "%Y-%m-%d").strftime("%b %Y")
    except Exception:
        return ""


HR = ("recruit", "talent", "hiring", "human resources", " hr", "hr ", "acquisition", "people")


def find_matches(conns, company, title="", terms=None):
    c = norm(company)
    kw = [w for w in norm(title).split() if len(w) > 3]
    terms = [norm(t) for t in (terms or []) if norm(t)]
    a, al, b, d = [], [], [], []
    for p in conns:
        pc, pos = norm(p["company"]), norm(p["position"])
        if c and pc and (c in pc or pc in c):
            a.append(dict(p, why="Works at " + company))
        elif any(t in pc or t in pos for t in terms):
            al.append(dict(p, why="Alumni / ex-colleague network"))
        elif any(t in " " + pos + " " for t in HR):
            b.append(dict(p, why="Recruiter / HR"))
        elif kw and sum(w in pos for w in kw) >= 2:
            d.append(dict(p, why="Similar role"))
    return a[:5] + al[:3] + b[:3] + d[:2]


def li_draft(p, job, hist, prior):
    P = profile()
    first = p["first"] or "there"
    co = job.get("company", "the company")
    title = job.get("title", "this role")
    y = re.search(r"\d{4}", p.get("connected_on", ""))
    sup = ", ".join((job.get("supported") or job.get("matched") or [])[:4]) or "cash management and transaction banking"
    at_co = bool(norm(p["company"]) and norm(co) and (norm(co) in norm(p["company"]) or norm(p["company"]) in norm(co)))
    proj = next((x["url"] for x in P.get("projects", []) if x.get("url")), "")
    facts = (f"They work as {p['position'] or 'a professional'} at {p['company'] or 'their company'}. "
             + (f"Connected since {y.group()}. " if y else "") + (f"Last conversation with them: {prior}. " if prior else ""))
    prompt = (f"Write a LinkedIn direct message, max 75 words, from {P['name']} to {first}. Use ONLY these facts: {facts}"
              f"Target role: {title} at {co}. Candidate's verified skills for this role: {sup}. "
              f"Candidate's note: {hist or 'none'}. Goal: politely ask whether they can refer him or help him get this role in any way "
              f"(a referral, an intro to the hiring manager, or advice). "
              f"{'Mention that they work at the same company.' if at_co else 'They may not work at the company; ask for an intro or advice.'} "
              f"Warm, specific, not pushy. Plain text, no subject line, sign with first name only. Invent nothing.")
    try:
        t = AG.llm_text(prompt).strip()
    except Exception:
        t = (f"Hi {first}, hope you're doing well."
             + (f" It's been a while since we last spoke ({prior})." if prior else (f" We've been connected since {y.group()}." if y else ""))
             + (" " + hist.strip() if hist.strip() else "")
             + (f" I saw that {co} has an opening for {title}, which lines up with my background in {sup}. "
                f"Since you're there, would you be open to referring me or sharing how the team hires?" if at_co else
                f" I'm exploring {title} at {co} and thought you might know the right person to speak with. "
                f"Could you introduce me or share any advice? My background: {sup}.")
             + " Thank you!\n\n" + P["name"].split()[0])
    return t + (f"\n\nMy work: {proj}" if proj else "")


# ---------------- resume builder ----------------
def build_resume():
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    base = G["find_resume"]()
    if not base or not base.lower().endswith(".docx"):
        raise HTTPException(400, "No .docx resume found. Upload your resume (.docx) in the Resume tab first.")
    P = profile()
    d = Document(base)
    pars = d.paragraphs
    full = "\n".join(p.text for p in pars).lower()
    changes = []
    blog = P["blog"].replace("https://", "").replace("http://", "")
    for p in pars:
        t = p.text.lower()
        if "@" in t and "linkedin" in t and p.runs:
            last = p.runs[-1]
            add = []
            li = P["linkedin"].replace("https://", "").replace("http://", "").replace("www.", "")
            if P["email"].lower() not in t:
                add.append(P["email"])
            if li.lower() not in t.replace("www.", ""):
                add.append(li)

            def clone(text, brk=False, p=p, last=last):
                r = p.add_run()
                if brk:
                    r.add_break()
                r.add_text(text)
                r.font.size = last.font.size
                r.font.name = last.font.name
                if last.font.color is not None and last.font.color.rgb is not None:
                    r.font.color.rgb = last.font.color.rgb

            if add:
                clone(" | " + " | ".join(add))
                changes.append("Added missing email/LinkedIn to header")
            if blog.lower() not in t:
                clone("Blog: " + blog, brk=True)
                changes.append("Added blog link to header")
            break
    host = lambda u: re.sub(r"^https?://|/.*$", "", u.lower())
    missing = [x for x in P.get("projects", []) if x.get("url") and host(x["url"]) not in full]

    def after(par, text, bold=False, italic=False, size=10.5, before=4):
        e = OxmlElement("w:p")
        par._p.addnext(e)
        n = Paragraph(e, par._parent)
        r = n.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        r.font.name = "Calibri"
        n.paragraph_format.space_before = Pt(before)
        n.paragraph_format.space_after = Pt(1)
        return n

    if missing:
        head = next((p for p in pars if p.text.strip().upper().startswith(
            ("TRANSACTION BANKING PROTOTYPES", "PROJECTS", "TRANSACTION BANKING PROJECTS"))), None)
        if head is None:
            head = after(pars[-1], "PROJECTS", bold=True, size=11, before=10)
            changes.append("Created PROJECTS section")
        blocks = []
        for x in missing:
            blocks += [(x["name"] + " | Live", dict(bold=True)), ("Link: " + x["url"], dict(italic=True, before=1))]
            if x.get("summary"):
                blocks.append((x["summary"], dict(size=9.5, before=1)))
        for text, kw in reversed(blocks):
            after(head, text, **kw)
        changes.append("Added %d project(s) from your profile" % len(missing))
    try:
        for rel in d.part.rels.values():
            if rel.is_external and "google.com/search?q=" in rel.target_ref:
                rel._target = rel.target_ref.split("q=", 1)[1]
                changes.append("Fixed redirect link")
    except Exception:
        pass
    os.makedirs("output", exist_ok=True)
    d.save(RESUME_OUT)
    return {"path": RESUME_OUT, "base": base, "changes": changes or ["Resume already contains all profile details"]}


# ---------------- API ----------------
class SearchReq(BaseModel):
    titles: str = ""
    location: str = "India"
    keywords: str = ""
    must: str = ""
    exclude: str = ""
    days: int = 0
    strict_location: bool = True
    min_score: int = 40
    max_jobs: int = 25
    max_queries: int = 10
    dry_run: bool = False
    only_live: bool = True
    exact_title: bool = True
    resume_path: str = ""


class Job(BaseModel):
    title: str
    company: str = ""
    location: str = ""
    url: str = ""
    portal: str = ""
    score: int = 0
    description: str = ""
    matched: List[str] = []
    gaps: List[str] = []
    supported: List[str] = []


class DraftReq(BaseModel):
    job: Job
    recruiter: str = "Hiring Team"
    history: str = ""
    extra: str = ""


class StatusReq(BaseModel):
    id: str
    status: str
    note: str = ""


class SendReq(BaseModel):
    to: str
    subject: str
    body: str
    app_id: str = ""


class LogReq(BaseModel):
    id: str = ""
    company: str = ""
    person: str
    message: str


class Csv(BaseModel):
    csv: str


class B64(BaseModel):
    b64: str


class KeyAdd(BaseModel):
    provider: str
    key: str
    label: str = ""


class KeyAct(BaseModel):
    provider: str
    id: str = ""
    action: str


def _load(f, d):
    try:
        return json.load(open(f))
    except Exception:
        return d


def install(app, g):
    global AG
    import careertrace_agent as ag
    AG = ag
    if not getattr(g["make_email"], "_pro", False):
        orig = g["make_email"]

        def signed(*a, **k):
            return sign(orig(*a, **k))
        signed._pro = True
        g["make_email"] = signed
    G.update(g)
    _o_text, _o_json = ag.llm_text, ag.llm_json

    def _t(pr):
        try:
            return llm_text(pr)
        except NoKey:
            return _o_text(pr)

    def _j(pr):
        try:
            return llm_json(pr)
        except NoKey:
            return _o_json(pr)
    ag.llm_text, ag.llm_json = _t, _j
    profile()
    r = APIRouter()
    state = STATE
    state["conns"] = _load("data/connections.json", [])
    state["msgs"] = _load("data/messages.json", {})

    @r.get("/pro", response_class=HTMLResponse)
    def page():
        return (STATIC / "pro.html").read_text(encoding="utf8")

    @r.get("/pro/app.js")
    def js():
        return Response((STATIC / "app.js").read_text(encoding="utf8"), media_type="application/javascript")

    @r.get("/api/pro/health")
    def health():
        return {"resume": G["find_resume"](), "resume_built": os.path.exists(RESUME_OUT),
                "connections": len(state["conns"]), "messages": len(state["msgs"]),
                "keys": {PROVIDERS[n]["label"]: (len(pool(n)) if enabled(n) else -1) for n in PROVIDERS},
                "last": STATE.get("last")}

    @r.post("/api/pro/search")
    def search(q: SearchReq):
        cfg = q.model_dump()
        res = do_search(cfg)
        s = settings()
        s["search"] = dict(cfg, dry_run=False)
        save_settings(s)
        return res

    @r.get("/api/pro/settings")
    def get_settings():
        return settings()

    @r.post("/api/pro/settings")
    def set_settings(p: dict):
        s = settings()
        if "autorun" in p:
            s["autorun"] = bool(p["autorun"])
        if p.get("provider") in PROVIDERS:
            s["providers_enabled"][p["provider"]] = bool(p.get("enabled", True))
        save_settings(s)
        return s

    @r.get("/api/pro/keys")
    def keys_list():
        d = _kload()
        return {"providers": [{"name": n, "label": m["label"], "kind": m["kind"], "hint": m["hint"], "enabled": enabled(n),
                               "keys": [{"id": e["id"], "masked": mask(e["key"]), "label": e.get("label", ""), "off": e.get("off", False),
                                         "status": e.get("status", ""), "used": e.get("used", "")} for e in d.get(n, [])]}
                              for n, m in PROVIDERS.items()]}

    @r.post("/api/pro/keys/add")
    def keys_add(b: KeyAdd):
        if b.provider not in PROVIDERS:
            raise HTTPException(400, "Unknown provider")
        k = re.sub(r"[\s\"']", "", b.key)
        if len(k) < 8:
            raise HTTPException(400, "That does not look like a key")
        if b.provider == "adzuna" and ":" not in k:
            raise HTTPException(400, "Adzuna format is APP_ID:APP_KEY")
        with _lock:
            d = _kload()
            if any(e["key"] == k for e in d[b.provider]):
                raise HTTPException(400, "That key is already added")
            d[b.provider].append({"id": uuid.uuid4().hex[:8], "key": k, "label": b.label.strip()[:30], "off": False, "status": "", "used": ""})
            _ksave(d)
        return {"ok": True}

    @r.post("/api/pro/keys/act")
    def keys_act(b: KeyAct):
        if b.provider not in PROVIDERS:
            raise HTTPException(400, "Unknown provider")
        with _lock:
            d = _kload()
            lst = d[b.provider]
            e = next((x for x in lst if x["id"] == b.id), None)
            if b.action == "reset":
                for x in lst:
                    x["status"] = ""
                _ksave(d)
                return {"ok": True, "detail": "Reset"}
            if not e:
                raise HTTPException(404, "Key not found")
            if b.action == "remove":
                lst.remove(e)
                _ksave(d)
                return {"ok": True, "detail": "Removed"}
            if b.action == "toggle":
                e["off"] = not e.get("off", False)
                _ksave(d)
                return {"ok": True, "detail": "Disabled" if e["off"] else "Enabled"}
            key = e["key"]
        if b.action == "test":
            ok, detail = test_key(b.provider, key)
            mark(b.provider, key, "" if ok else detail[:80])
            return {"ok": ok, "detail": detail}
        raise HTTPException(400, "Unknown action")

    @r.get("/api/pro/exports")
    def exps():
        out = []
        for f in sorted(Path("exports").glob("*"), key=lambda p: -p.stat().st_mtime):
            if f.suffix in (".xlsx", ".csv"):
                out.append({"name": f.name, "kb": round(f.stat().st_size / 1024, 1),
                            "time": dt.datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y-%m-%d %H:%M")})
        return out[:60]

    @r.get("/api/pro/exports/{name}")
    def expf(name: str):
        f = Path("exports") / Path(name).name
        if not f.exists():
            raise HTTPException(404, "No such file")
        return FileResponse(str(f), filename=f.name)

    @r.get("/api/pro/statuses")
    def statuses():
        return STATUSES

    @r.get("/api/pro/apps")
    def apps():
        return list_apps()

    @r.post("/api/pro/apps")
    def add(j: Job, status: str = "Saved"):
        row = add_app(j.model_dump(), status)
        if status != "Saved" and row["Status"] != status:
            set_status(row["Application ID"], status)
        return row

    @r.post("/api/pro/status")
    def set_st(s: StatusReq):
        return {"ok": set_status(s.id, s.status, s.note)}

    @r.post("/api/pro/email")
    def email(d: DraftReq):
        row = add_app(d.job.model_dump(), "Saved")
        P = profile()
        j = d.job
        try:
            txt = G["make_email"](j.title, j.company, j.location, j.supported or j.matched, j.gaps, d.extra)
        except Exception:
            txt = sign(f"Subject: Application: {j.title} - {P['name']}\n\nDear Hiring Team,\n\nI am writing to apply for the "
                       f"{j.title} role at {j.company}. My resume is attached and I would welcome a conversation.\n\nRegards,\n[Your Name]")
        m = re.match(r"\s*Subject:\s*(.*?)\n+(.*)", txt, re.S)
        subj, body = (m.group(1).strip(), m.group(2).strip()) if m else (f"Application: {j.title} - {P['name']}", txt)
        if d.recruiter and d.recruiter != "Hiring Team":
            body = body.replace("Hiring Team", d.recruiter, 1)
        return {"subject": subj, "body": body, "app_id": row["Application ID"]}

    @r.post("/api/pro/email/send")
    def send(s: SendReq):
        gk = pool("gmail")
        pw = gk[0]["key"] if gk else ""
        P = profile()
        if not pw:
            raise HTTPException(400, "Add a Gmail app password in Settings > API keys to send from the app, or use 'Open in Gmail'.")
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = P["email"], s.to, s.subject
        m.set_content(s.body)
        att = RESUME_OUT if os.path.exists(RESUME_OUT) else G["find_resume"]()
        if att and att.lower().endswith(".docx"):
            m.add_attachment(open(att, "rb").read(), maintype="application",
                             subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                             filename="Amit_Mankar_Resume.docx")
        try:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as sm:
                sm.login(P["email"], pw)
                sm.send_message(m)
        except Exception as e:
            raise HTTPException(400, f"Send failed: {e}")
        if s.app_id:
            set_status(s.app_id, "Recruiter Contacted", "Email sent to " + s.to, **{"Contact Email": s.to})
        return {"ok": True}

    @r.post("/api/pro/connections")
    def conns(c: Csv):
        state["conns"] = parse_connections(c.csv)
        json.dump(state["conns"], open("data/connections.json", "w"))
        return {"count": len(state["conns"])}

    @r.post("/api/pro/messages")
    def msgs(c: Csv):
        state["msgs"] = parse_messages(c.csv)
        json.dump(state["msgs"], open("data/messages.json", "w"))
        return {"count": len(state["msgs"])}

    @r.post("/api/pro/outreach")
    def outreach(d: DraftReq):
        ms = find_matches(state["conns"], d.job.company, d.job.title, split(profile().get("alumni_terms", "")))
        jd = d.job.model_dump()

        def one(p):
            prior = _fmt(state["msgs"].get((p.get("url") or "").lower().rstrip("/"), ""))
            return {"person": f"{p['first']} {p['last']}".strip(), "position": p["position"], "company": p["company"],
                    "profile": p["url"], "prior": prior, "why": p.get("why", ""), "message": li_draft(p, jd, d.history, prior)}

        with ThreadPoolExecutor(4) as ex:
            res = list(ex.map(one, ms))
        return {"connections_loaded": len(state["conns"]), "matches": res}

    @r.post("/api/pro/outreach/log")
    def outlog(l: LogReq):
        log_outreach(l.id, l.company, l.person, l.message)
        if l.id:
            set_status(l.id, "Referral Requested", "LinkedIn draft for " + l.person)
        return {"ok": True}

    @r.get("/api/pro/profile")
    def gp():
        return profile()

    @r.post("/api/pro/profile")
    def sp(p: dict):
        json.dump(p, open("data/profile.json", "w"), indent=1)
        return {"ok": True}

    @r.post("/api/pro/resume/upload")
    def rup(b: B64):
        os.makedirs("data", exist_ok=True)
        open("data/resume.docx", "wb").write(base64.b64decode(b.b64))
        return {"ok": True}

    @r.post("/api/pro/resume/build")
    def rb():
        return build_resume()

    @r.get("/api/pro/resume/download")
    def rd():
        p = RESUME_OUT if os.path.exists(RESUME_OUT) else G["find_resume"]()
        if not p:
            raise HTTPException(404, "No resume")
        return FileResponse(p, filename="Amit_Mankar_Resume.docx")

    @r.get("/api/pro/export")
    def export():
        _wb()
        return FileResponse(TRACKER, filename="CareerTrace_Tracker.xlsx")

    app.include_router(r)

    @app.middleware("http")
    async def inject(request, call_next):
        if request.method == "GET" and request.url.path == "/" and not request.query_params.get("classic"):
            return RedirectResponse("/pro")
        resp = await call_next(request)
        p = request.url.path
        if request.method == "GET" and not p.startswith(("/pro", "/api", "/docs", "/redoc", "/openapi")) \
                and "text/html" in resp.headers.get("content-type", ""):
            body = b"".join([c async for c in resp.body_iterator]).decode("utf-8", "ignore")
            tag = '<script src="/pro/app.js"></script>'
            body = body.replace("</body>", tag + "</body>", 1) if "</body>" in body else body + tag
            h = {k: v for k, v in resp.headers.items() if k.lower() not in ("content-length", "content-type")}
            return Response(body, status_code=resp.status_code, headers=h, media_type="text/html")
        return resp

    if os.environ.get("CT_AUTORUN") == "1":
        threading.Thread(target=autorun, daemon=True).start()
    print("CareerTrace Pro loaded: open the app and click the blue Pro button (or /pro).")
